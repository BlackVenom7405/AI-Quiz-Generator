import os
import fitz  # PyMuPDF
import docx

def parse_pdf(file_path):
    """
    Extract text from a PDF file using PyMuPDF (fitz).
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    doc = fitz.open(file_path)
    text_content = []
    
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text_content.append(page.get_text())
        
    doc.close()
    extracted_text = "\n".join(text_content).strip()
    
    if not extracted_text:
        raise ValueError("Could not extract any text from the PDF file. It might be empty or scanned images.")
        
    return extracted_text

def parse_docx(file_path):
    """
    Extract text from a DOCX file using python-docx.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    doc = docx.Document(file_path)
    text_content = []
    
    # Extract from paragraphs
    for para in doc.paragraphs:
        if para.text:
            text_content.append(para.text)
            
    # Extract from tables (optional, but good for completeness)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    text_content.append(cell.text)
                    
    extracted_text = "\n".join(text_content).strip()
    
    if not extracted_text:
        raise ValueError("Could not extract any text from the DOCX file. It might be empty.")
        
    return extracted_text

def allowed_file(filename, allowed_extensions):
    """
    Check if the file extension is allowed.
    """
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in allowed_extensions
